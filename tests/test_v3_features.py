"""Проверки третьей версии: города с координатами, гонка зеркал Overpass и кэш,
участники и удаление, рассылка с форматированием, экспорт из 7 колонок."""

import asyncio
import base64
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

import api
import countries
import db
import export
import parser as search
from parser import GeoError


class CountryCityTests(unittest.TestCase):
    def test_every_country_has_60_to_150_cities_sorted_by_population(self):
        for code in countries.ORDER:
            cities = countries.cities(code)
            self.assertGreaterEqual(len(cities), 60, code)
            self.assertLessEqual(len(cities), 150, code)
            self.assertEqual(len(cities), len(set(cities)), code)
        self.assertEqual(countries.cities("RU")[:2], ["Москва", "Санкт-Петербург"])
        self.assertEqual(countries.cities("PL")[0], "Варшава")

    def test_city_bbox_without_geocoder(self):
        south, west, north, east = countries.city_bbox("RU", "Москва")
        self.assertTrue(south < 55.75 < north and west < 37.62 < east)
        # маленький город — маленький район, расширение — больше
        small = countries.city_bbox("RU", "Арзамас") or countries.city_bbox("RU", countries.cities("RU")[-1])
        wide = countries.city_bbox("RU", "Москва", scale=2.2)
        self.assertLess(small[2] - small[0], north - south)
        self.assertGreater(wide[2] - wide[0], north - south)
        # старые названия и местные названия тоже находятся
        self.assertIsNotNone(countries.city_bbox("PL", "Warszawa"))
        self.assertIsNotNone(countries.city_bbox("UA", "Kyiv"))
        self.assertIsNotNone(countries.city_bbox("MD", "Кишинев"))
        self.assertIsNone(countries.city_bbox("RU", "Несуществующийград"))

    def test_quick_query_asks_only_for_businesses_with_phone(self):
        query = search.build_query((1, 2, 3, 4), 0, "", country="RU", require_phone=True)
        self.assertIn("phone|contact:phone|mobile|contact:mobile", query)
        self.assertNotIn("phone|contact", search.build_query((1, 2, 3, 4), 0, "", country="RU"))


class OverpassRaceTests(unittest.IsolatedAsyncioTestCase):
    async def test_fastest_mirror_wins_and_result_is_cached(self):
        calls = []

        async def one(_session, endpoint, _query, delay):
            calls.append(endpoint)
            if "slow" in endpoint:
                await asyncio.sleep(5)
                return [{"id": 1}]
            await asyncio.sleep(0.01)
            return [{"id": 2}]

        store = {}

        async def cache_get(query):
            return store.get(query)

        async def cache_set(query, elements):
            store[query] = elements

        with patch("parser._overpass_one", one), patch("parser.OVERPASS_URL", "https://slow.test"), patch(
            "parser.OVERPASS_FALLBACK_URLS", ("https://fast.test",)
        ), patch("parser.CACHE_GET", cache_get), patch("parser.CACHE_SET", cache_set):
            started = asyncio.get_running_loop().time()
            rows = await search.overpass(None, "Q")
            self.assertLess(asyncio.get_running_loop().time() - started, 2)
            self.assertEqual(rows, [{"id": 2}])
            self.assertEqual(store["Q"], [{"id": 2}])
            calls.clear()
            self.assertEqual(await search.overpass(None, "Q"), [{"id": 2}])
            self.assertEqual(calls, [])  # из кэша, без запросов

    async def test_all_mirrors_down_gives_clear_error(self):
        async def one(*_args):
            raise search._Retryable(504)

        with patch("parser._overpass_one", one), patch("parser.CACHE_GET", None), patch("parser.OVERPASS_ROUND_PAUSE", 0):
            with self.assertRaises(GeoError) as ctx:
                await search.overpass(None, "Q")
        self.assertIn("504", str(ctx.exception))


class TempDB(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="parsercc-v3-")
        self.original = db.DB_PATH
        db.DB_PATH = Path(self.temp.name) / "t.sqlite"
        await db.init()

    async def asyncTearDown(self):
        db.DB_PATH = self.original
        self.temp.cleanup()


class DbV3Tests(TempDB):
    async def test_osm_cache_roundtrip_and_ttl(self):
        await db.osm_cache_set("query", [{"id": 1, "tags": {"name": "Тест"}}])
        self.assertEqual(await db.osm_cache_get("query"), [{"id": 1, "tags": {"name": "Тест"}}])
        self.assertIsNone(await db.osm_cache_get("query", ttl=-1))
        self.assertIsNone(await db.osm_cache_get("other"))

    async def test_members_and_delete(self):
        await db.upsert_user(10, "anna", "Анна")
        await db.mark_bot_started(10)
        await db.upsert_user(11, "boris", "Борис")
        await db.mark_bot_started(11)
        await db.mark_app_opened(11)
        members = {m["id"]: m for m in await db.members()}
        self.assertFalse(members[10]["app_opened"])
        self.assertTrue(members[11]["app_opened"])
        self.assertEqual([m["id"] for m in await db.members("bor")], [11])
        self.assertTrue(await db.delete_user(10))
        user = await db.get_user(10)
        self.assertEqual((user["banned"], user["deleted"]), (1, 1))
        self.assertNotIn(10, [m["id"] for m in await db.members()])
        self.assertNotIn(10, [u["id"] for u in await db.all_users_stats()])
        self.assertNotIn(10, await db.approved_ids())
        # разбан возвращает участника
        await db.set_banned(10, False)
        self.assertIn(10, [u["id"] for u in await db.all_users_stats()])


class ExportTests(unittest.TestCase):
    def test_only_seven_columns_and_yandex_link(self):
        lead = {
            "name": "Барбер Топ", "niche": "Барбершопы", "country": "RU", "city": "Москва",
            "c": {"address": "Тверская, 1", "map": "https://yandex.ru/maps/?x"}, "phone": "+79261112233",
            "status": "written", "info": {"phones": ["+7 926 111-22-33", "+7 495 000-00-00"], "lat": 55.75, "lon": 37.61},
        }
        sheet = load_workbook(io.BytesIO(export.build_xlsx([lead]))).active
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            ["Название", "Ниша", "Страна", "Город", "Адрес", "Номер телефона", "Статус"],
        )
        row = [cell.value for cell in sheet[2]]
        self.assertEqual(row[:4], ["Барбер Топ", "Барбершопы", "Россия", "Москва"])
        self.assertEqual(row[4], "Тверская, 1")
        self.assertEqual(row[5], "+7 926 111-22-33, +7 495 000-00-00")
        self.assertEqual(row[6], "Написал")
        link = sheet.cell(row=2, column=5).hyperlink.target
        self.assertTrue(link.startswith("https://yandex.ru/maps/?"))
        self.assertIn("37.61%2C55.75", link)


class TelegramHtmlTests(unittest.TestCase):
    def test_only_telegram_tags_survive(self):
        text = api.telegram_html('<b>Жирный</b> <i>курсив</i> <script>x</script> <a href="https://t.me/x">ссылка</a> 5 < 6 & 7')
        self.assertIn("<b>Жирный</b>", text)
        self.assertIn("<i>курсив</i>", text)
        self.assertIn('<a href="https://t.me/x">ссылка</a>', text)
        self.assertIn("&lt;script&gt;", text)
        self.assertIn("5 &lt; 6 &amp; 7", text)
        self.assertEqual(api.telegram_html("<b>без конца"), "<b>без конца</b>")


class AdminApiV3Tests(unittest.TestCase):
    def setUp(self):
        from fastapi.testclient import TestClient

        self.temp = tempfile.TemporaryDirectory(prefix="parsercc-api3-")
        self.original = db.DB_PATH
        db.DB_PATH = Path(self.temp.name) / "api.sqlite"
        self.user = {"id": 1, "username": "owner", "first_name": "O"}
        self.broadcast = AsyncMock(return_value=(1, 1))
        self.patchers = [
            patch("api.verify_init_data", side_effect=lambda _d: dict(self.user)),
            patch("config.ADMIN_IDS", {1}),
            patch("api._notify", AsyncMock()),
            patch("bot.broadcast_to_approved", self.broadcast),
        ]
        for item in self.patchers:
            item.start()
        self.client_cm = __import__("fastapi.testclient", fromlist=["TestClient"]).TestClient(api.app)
        self.client = self.client_cm.__enter__()
        asyncio.run(db.upsert_user(1, "owner", "O"))

    def tearDown(self):
        self.client_cm.__exit__(None, None, None)
        for item in reversed(self.patchers):
            item.stop()
        db.DB_PATH = self.original
        self.temp.cleanup()

    def test_members_delete_and_broadcast(self):
        h = {"X-Init-Data": "x"}
        self.user = {"id": 55, "username": "client", "first_name": "C"}
        me = self.client.get("/api/me", headers=h).json()
        self.assertEqual(me["id"], 55)
        self.user = {"id": 1, "username": "owner", "first_name": "O"}
        asyncio.run(db.set_approved(1, 1))
        with patch("api.is_admin", lambda u: u["id"] == 1):
            members = self.client.get("/api/admin/members", headers=h).json()
            self.assertIn(55, [m["id"] for m in members["members"] if m["app_opened"]])
            self.assertEqual(self.client.delete("/api/admin/users/55", headers=h).status_code, 200)
            self.assertNotIn(55, [m["id"] for m in self.client.get("/api/admin/members", headers=h).json()["members"]])
            # удалённый участник получает бан
            self.user = {"id": 55, "username": "client", "first_name": "C"}
            banned = self.client.get("/api/me", headers=h)
            self.assertEqual(banned.status_code, 403)
            self.assertIn("Вы забанены, свяжитесь с @SUN9ISE", banned.json()["detail"]["message"])
            self.user = {"id": 1, "username": "owner", "first_name": "O"}
            pixel = base64.b64encode(
                bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                              "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082")
            ).decode()
            response = self.client.post("/api/admin/broadcast", headers=h, json={
                "text": "<b>Привет</b> <i>всем</i>", "image": "data:image/png;base64," + pixel,
            })
            self.assertEqual(response.status_code, 200, response.text)
            bad = self.client.post("/api/admin/broadcast", headers=h, json={"text": "", "image": "data:text/html;base64,AA=="})
            self.assertEqual(bad.status_code, 422)
        import time
        for _ in range(50):
            if self.broadcast.await_count:
                break
            time.sleep(0.02)
        args = self.broadcast.await_args
        self.assertEqual(args.args[0], "<b>Привет</b> <i>всем</i>")
        self.assertTrue(args.kwargs["image"].startswith(b"\x89PNG"))


if __name__ == "__main__":
    unittest.main()


class AdminConfigTests(unittest.TestCase):
    def test_several_admins_ids_and_usernames(self):
        import importlib, os
        import config
        old = os.environ.get("ADMIN_IDS")
        try:
            os.environ["ADMIN_IDS"] = "123456789; @SUN9ISE  987654321, мусор"
            cfg = importlib.reload(config)
            self.assertEqual(cfg.ADMIN_IDS, {123456789, 987654321})
            self.assertEqual(cfg.ADMIN_USERNAMES, {"sun9ise"})
            self.assertTrue(cfg.is_admin_user(5, "Sun9ise"))
            self.assertTrue(cfg.is_admin_user(987654321))
            self.assertFalse(cfg.is_admin_user(5, "other"))
            cfg.validate_config.__globals__["BOT_TOKEN"] = "x"
        finally:
            if old is None:
                os.environ.pop("ADMIN_IDS", None)
            else:
                os.environ["ADMIN_IDS"] = old
            importlib.reload(config)


class AdminByUsernameTests(TempDB):
    async def test_username_admin_gets_admin_row(self):
        with patch("db.is_admin_user", lambda uid, name=None: (name or "").lower() == "sun9ise"):
            row = await db.upsert_user(77, "SUN9ISE", "Alex")
        self.assertEqual((row["is_admin"], row["approved"]), (1, 1))
        self.assertIn(77, await db.admin_ids())


class PictureFallbackTests(TempDB):
    async def test_trial_over_has_close_picture_and_falls_back_to_text(self):
        import bot
        from aiogram.exceptions import TelegramBadRequest

        with patch("bot.bot.send_photo", AsyncMock(return_value=SimpleNamespace(photo=[SimpleNamespace(file_id="F")]))) as photo:
            await bot.send_trial_over(10, 10)
        self.assertTrue(str(photo.await_args.args[1].path).endswith("access_closed.jpg"))
        self.assertIn("<b>Пробный доступ закончился</b>", photo.await_args.kwargs["caption"])
        self.assertEqual(photo.await_args.kwargs["reply_markup"].inline_keyboard[0][0].text, "Вступить в команду")
        bad = TelegramBadRequest(method=SimpleNamespace(), message="wrong file identifier")
        with patch("bot.bot.send_photo", AsyncMock(side_effect=bad)), patch("bot.bot.send_message", AsyncMock()) as text:
            await bot.send_trial_over(10, 10)
        self.assertIn("Пробный доступ закончился", text.await_args.args[1])


class ReviewFixTests(unittest.TestCase):
    def test_broadcast_html_is_always_valid(self):
        self.assertEqual(api.telegram_html("a </b> b"), "a &lt;/b&gt; b")
        self.assertEqual(api.telegram_html("<b><i>x</b>"), "<b><i>x&lt;/b&gt;</i></b>")
        self.assertEqual(api.telegram_html('<a href="https://x.ru/?a=1&b=2">l</a>'), '<a href="https://x.ru/?a=1&amp;b=2">l</a>')

    def test_lead_word(self):
        self.assertEqual([api.lead_word(n) for n in (1, 3, 5, 11, 21)], ["лид", "лида", "лидов", "лидов", "лид"])

    def test_bot_number_parsing(self):
        import bot
        self.assertIsNone(bot._num("--5", signed=True))
        self.assertIsNone(bot._num("²"))
        self.assertEqual(bot._num("-5", signed=True), -5)


class GoogleFirstTests(unittest.IsolatedAsyncioTestCase):
    async def test_google_first_and_osm_skipped_when_enough(self):
        places = [{
            "id": f"P{i}", "displayName": {"text": f"Барбер {i}"}, "businessStatus": "OPERATIONAL",
            "internationalPhoneNumber": f"+7 926 300-{i:02d}-{i:02d}", "googleMapsUri": f"https://maps.google.com/?cid={i}",
            "location": {"latitude": 55.75, "longitude": 37.6 + i / 1000}, "shortFormattedAddress": f"Тверская, {i}",
        } for i in range(1, 40)]
        calls = []

        async def text_search(session, key, text, country, **kw):
            calls.append(kw.get("bbox"))
            return places

        overpass = AsyncMock(return_value=[])
        with patch("parser.GOOGLE_PLACES_API_KEY", "k"), patch("google_places.text_search", text_search), \
                patch("parser.overpass", overpass), patch("parser.GEOAPIFY_API_KEY", ""), patch("parser.ANALYZE_SITES", False):
            result = await search.run_search("RU", "Москва", 0, "", "all", 10)
        self.assertEqual(len(result), 10)
        self.assertIsNotNone(calls[0])  # поиск ограничен районом города
        overpass.assert_not_awaited()  # Google хватило — карты OSM не дёргали
